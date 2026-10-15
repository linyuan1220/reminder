#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
任务提醒（Windows）
读取同目录下的“提醒任务.xlsx”提醒任务，按周期在指定时间弹窗强提醒。
支持周期：每天 / 每周(周几) / 每月(几号) / 每年(几月几号) / 隔N天 / 指定日期（一次性）。
强/弱/错过 三种提醒可分别配置不同的提示音（可开关）。
错过的提醒（未开机时到期的）会用黄色闪动窗口补提醒。
本程序为纯离线工具：代码层面已禁止网络连接，不与外部进行任何通信。
"""

import os
import sys
import json
import stat
import time
import queue
import threading
import hashlib
import struct
import math
import wave
import datetime as dt
import calendar
import winreg

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

# ---------- 第三方库（打包时会一并带上） ----------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw
import pystray
import pygame   # 提示音播放引擎（wav/mp3/ogg，循环可靠）


APP_NAME = "任务提醒"

# =====================================================================
# 路径
# =====================================================================
def app_dir():
    """可执行文件(.exe)或脚本所在目录。"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = app_dir()
EXCEL_PATH = os.path.join(BASE_DIR, "提醒任务.xlsx")
STATE_PATH = os.path.join(BASE_DIR, "reminder_state.json")
SETTINGS_PATH = os.path.join(BASE_DIR, "settings.json")

DEFAULT_SOUND_NAMES = {
    "strong": "音效_强提醒.wav",
    "weak": "音效_弱提醒.wav",
    "missed": "音效_错过提醒.wav",
}

# =====================================================================
# 网络限制（纯离线）
# =====================================================================
def block_network():
    """代码层面禁止本程序发起任何网络连接，不与外部通信。"""
    def _deny(*a, **k):
        raise OSError("网络访问已禁用：本程序为纯离线工具，不与外部通信")
    import socket
    class _OfflineSocket(socket.socket):
        def connect(self, address):
            _deny()
        def sendto(self, *a, **k):
            _deny()
    try:
        socket.socket = _OfflineSocket
        socket.create_connection = _deny
    except Exception:
        pass


# =====================================================================
# 日期 / 时间 / 周期字段解析
# =====================================================================
DEFAULT_TIME = dt.time(9, 0)
WEEKDAY_CN = {1: "周一", 2: "周二", 3: "周三", 4: "周四", 5: "周五", 6: "周六", 7: "周日"}


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
        except Exception:
            return None
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y-%m"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_time(v):
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.time):
        return v
    if isinstance(v, (int, float)):
        total = float(v) * 86400
        h = int(total // 3600) % 24
        m = int((total % 3600) // 60)
        return dt.time(h, m)
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%H点%M分"):
        try:
            return dt.datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return None


_CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "日": 7, "天": 7, "七": 7}


def parse_weekday(v):
    if isinstance(v, (int, float)):
        n = int(v)
        return n if 1 <= n <= 7 else None
    s = str(v).strip()
    if not s:
        return None
    if s.isdigit():
        n = int(s)
        return n if 1 <= n <= 7 else None
    body = s
    for p in ("周", "星期", "礼拜"):
        body = body.replace(p, "")
    if body in _CN_NUM:
        return _CN_NUM[body]
    return None


def parse_day_of_month(v):
    if isinstance(v, (int, float)):
        n = int(v)
        return n if 1 <= n <= 31 else None
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("号", "").replace("日", "").strip()
    if s.isdigit():
        n = int(s)
        return n if 1 <= n <= 31 else None
    return None


def parse_month_day(v):
    if isinstance(v, dt.datetime):
        return (v.month, v.day)
    if isinstance(v, dt.date):
        return (v.month, v.day)
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("月", "-").replace("号", "").replace("日", "")
    for sep in ("-", "/", ".", " ", ","):
        if sep in s:
            parts = [p for p in s.split(sep) if p != ""]
            if len(parts) == 2:
                try:
                    m, d = int(parts[0]), int(parts[1])
                    if 1 <= m <= 12 and 1 <= d <= 31:
                        return (m, d)
                except ValueError:
                    return None
    if s.isdigit() and len(s) in (3, 4):
        if len(s) == 3:
            m, d = int(s[0]), int(s[1:])
        else:
            m, d = int(s[:2]), int(s[2:])
        if 1 <= m <= 12 and 1 <= d <= 31:
            return (m, d)
    return None


def clamp_day(year, month, day):
    last = calendar.monthrange(year, month)[1]
    return min(day, last)


def advance_month(y, m):
    m += 1
    if m > 12:
        m, y = 1, y + 1
    return y, m


def _step(cand, after, step, strict):
    if strict:
        while cand <= after:
            cand += step
    else:
        while cand < after:
            cand += step
    return cand


def next_occurrence(task, after, strict=False):
    """计算最近一次提醒时间（>= after；strict 时严格 > after）。"""
    after = after.replace(second=0, microsecond=0)
    t = task["time"] or DEFAULT_TIME
    cycle = (task["cycle"] or "").strip()

    if cycle in ("每天", "每日", "天天"):
        cand = dt.datetime.combine(after.date(), t)
        return _step(cand, after, dt.timedelta(days=1), strict)

    if cycle in ("每周", "每周一次"):
        wd = task.get("weekday")
        if not wd:
            return None
        days_ahead = (wd - after.isoweekday()) % 7
        cand = dt.datetime.combine(after.date() + dt.timedelta(days=days_ahead), t)
        return _step(cand, after, dt.timedelta(days=7), strict)

    if cycle in ("每月", "每月一次"):
        dom = task.get("day_of_month")
        if not dom:
            return None
        y, m = after.year, after.month
        while True:
            day = clamp_day(y, m, dom)
            cand = dt.datetime(y, m, day, t.hour, t.minute)
            if (cand > after) if strict else (cand >= after):
                return cand
            y, m = advance_month(y, m)

    if cycle in ("每年", "每年一次"):
        md = task.get("month_day")
        if not md:
            return None
        mo, da = md
        y = after.year
        while True:
            day = clamp_day(y, mo, da)
            cand = dt.datetime(y, mo, day, t.hour, t.minute)
            if (cand > after) if strict else (cand >= after):
                return cand
            y += 1

    if cycle in ("隔N天", "隔几天", "自定义天数", "间隔"):
        try:
            n = int(task.get("num") or 1)
        except (TypeError, ValueError):
            n = 1
        if n < 1:
            n = 1
        start = task.get("start_date")
        base = start if start is not None else after.date()
        cand = dt.datetime.combine(base, t)
        return _step(cand, after, dt.timedelta(days=n), strict)

    if cycle in ("指定日期", "一次性", "某一天", "单次", "仅一次"):
        d = task.get("one_date")
        if d is None:
            return None
        return dt.datetime.combine(d, t)

    cand = dt.datetime.combine(after.date(), t)
    return _step(cand, after, dt.timedelta(days=1), strict)


# =====================================================================
# 提示音（音效）
# =====================================================================
def default_sound_path(key):
    return os.path.join(BASE_DIR, DEFAULT_SOUND_NAMES.get(key, "音效_强提醒.wav"))


def _write_tone_wav(path, seq):
    """按 (频率Hz, 时长秒, 幅度) 序列生成一个 wav 文件。"""
    sr = 22050
    frames = []
    for freq, dur, amp in seq:
        n = int(sr * dur)
        for i in range(n):
            env = min(1.0, i / (sr * 0.01), (n - i) / (sr * 0.03))
            val = amp * env * math.sin(2 * math.pi * freq * i / sr)
            frames.append(int(max(-1, min(1, val)) * 32767))
    try:
        with wave.open(path, "w") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(sr)
            w.writeframes(struct.pack("<%dh" % len(frames), *frames))
        return True
    except Exception:
        return False


def ensure_default_sounds():
    """生成三种默认音效（强/弱/错过），互不相同。"""
    melodies = {
        "strong": [(700, 0.22, 0.55), (900, 0.35, 0.65)],
        "weak": [(660, 0.20, 0.42)],
        "missed": [(520, 0.20, 0.5), (740, 0.20, 0.5), (740, 0.32, 0.6)],
    }
    for key, fname in DEFAULT_SOUND_NAMES.items():
        p = os.path.join(BASE_DIR, fname)
        if not os.path.exists(p):
            _write_tone_wav(p, melodies[key])


def load_settings():
    defaults = {"sound_enabled": True,
                "sound_strong": "", "sound_weak": "", "sound_missed": ""}
    try:
        with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            for k in defaults:
                if k in data:
                    defaults[k] = data[k]
    except Exception:
        pass
    return defaults


def save_settings(settings):
    tmp = SETTINGS_PATH + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
        os.replace(tmp, SETTINGS_PATH)
    except Exception:
        pass


class SoundPlayer:
    """循环播放提示音，直到 stop()。
    优先用 pygame（对 wav/mp3/ogg 播放与循环最可靠）；
    失败则退回 Windows MCI；再失败退回循环蜂鸣。"""

    def __init__(self):
        self._kind = None            # 'pygame' / 'mci' / 'beep'
        self._pygame_sound = None
        self._mixer_ready = False
        self._mci_alias = None
        self._mci_thread = None
        self._stop_event = threading.Event()

    def _ensure_mixer(self):
        if self._mixer_ready:
            return True
        try:
            if not pygame.mixer.get_init():
                pygame.mixer.init(frequency=44100)
            self._mixer_ready = True
        except Exception:
            self._mixer_ready = False
        return self._mixer_ready

    def play_loop(self, path):
        self.stop()
        self._stop_event = threading.Event()
        if path and os.path.exists(path):
            # 1) pygame（wav/mp3/ogg，循环最可靠）
            if self._ensure_mixer():
                try:
                    snd = pygame.mixer.Sound(path)
                    snd.play(loops=-1)
                    self._kind = "pygame"
                    self._pygame_sound = snd
                    return
                except Exception:
                    pass
            # 2) Windows MCI（其它格式，尽力而为）
            if self._mci_start(path):
                self._kind = "mci"
                self._mci_watch()
                return
        # 3) 兜底：循环蜂鸣
        self._start_beep_loop()

    # ---------- Windows MCI（其它格式） ----------
    def _mci_start(self, path):
        try:
            import ctypes
            winmm = ctypes.windll.winmm
            alias = "trem" + str(int(time.time() * 1000))[-9:]
            self._mci_alias = alias
            r = winmm.mciSendStringW(f'open "{path}" alias {alias}', None, 0, 0)
            if r != 0:
                return False
            r = winmm.mciSendStringW(f"play {alias} from 0", None, 0, 0)
            return r == 0
        except Exception:
            return False

    def _mci_watch(self):
        # 轮询播放状态，播完自动从头再播，实现循环
        def watch():
            try:
                import ctypes
                winmm = ctypes.windll.winmm
                alias = self._mci_alias
                buf = ctypes.create_unicode_buffer(128)
                while not self._stop_event.is_set() and self._kind == "mci":
                    r = winmm.mciSendStringW(f"status {alias} mode", buf, 128, 0)
                    if r != 0:
                        break
                    if buf.value.strip().lower() != "playing":
                        winmm.mciSendStringW(f"play {alias} from 0", None, 0, 0)
                    self._stop_event.wait(0.25)
            except Exception:
                pass
        self._mci_thread = threading.Thread(target=watch, daemon=True)
        self._mci_thread.start()

    # ---------- 兜底：循环蜂鸣 ----------
    def _start_beep_loop(self):
        def loop():
            try:
                import winsound
                while not self._stop_event.is_set():
                    winsound.Beep(880, 220)
                    self._stop_event.wait(0.45)
            except Exception:
                pass
        threading.Thread(target=loop, daemon=True).start()
        self._kind = "beep"

    # ---------- 停止 ----------
    def stop(self):
        self._stop_event.set()
        try:
            if self._kind == "pygame" and self._pygame_sound is not None:
                self._pygame_sound.stop()
        except Exception:
            pass
        try:
            if self._kind == "mci" and self._mci_alias:
                import ctypes
                ctypes.windll.winmm.mciSendStringW(
                    f"close {self._mci_alias}", None, 0, 0)
        except Exception:
            pass
        self._kind = None
        self._pygame_sound = None
        self._mci_alias = None


# =====================================================================
# Excel 读写
# =====================================================================
HEADERS = ["提醒内容", "周期类型", "周期数值(天)", "周几(每周)",
           "每月几号(每月)", "几月几号(每年)", "起始/指定日期", "提醒时间(时:分)", "提醒强度(强/弱)"]

EXAMPLES = [
    ["喝水", "隔N天", 2, "", "", "", "2026-10-13", "10:00", "弱"],
    ["每周例会", "每周", "", "周一", "", "", "", "14:30", "强"],
    ["发工资", "每月", "", "", "15", "", "", "09:00", "弱"],
    ["缴物业费", "每年", "", "", "", "10-01", "", "09:00", "弱"],
    ["每天午休", "每天", "", "", "", "", "", "12:30", "弱"],
    ["交房租", "指定日期", "", "", "", "", "2026-10-20", "18:00", "强"],
]


def create_template(path):
    """生成带标题、表头和下拉选择(数据验证)的 Excel 模板，单元格放大。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "提醒任务"

    n_cols = len(HEADERS)          # 9 列 A~I
    last_col = openpyxl.utils.get_column_letter(n_cols)

    # ---------- 第1行：标题 ----------
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = "任务提醒 · 提醒任务表"
    c.font = Font(size=16, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="2F5BB3")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # ---------- 第2行：表头 ----------
    ws.append(HEADERS)             # 落在第2行
    thin = Side(style="thin", color="9EB3D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5A86D6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[2].height = 26

    # ---------- 示例数据（第3行起）----------
    for row in EXAMPLES:
        ws.append(row)
    for r in range(3, 3 + len(EXAMPLES)):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[r].height = 26

    # 加大默认行高与列宽
    ws.sheet_format.defaultRowHeight = 24
    widths = [26, 12, 13, 11, 13, 13, 15, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # ---------- 下拉选择（数据验证，从第3行起）----------
    from openpyxl.worksheet.datavalidation import DataValidation
    n_rows = 500

    def add_dv(column_letter, formula, error="", prompt="", title=""):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        if title:
            dv.errorTitle = title
        if error:
            dv.error = error
        if prompt:
            dv.promptTitle = title or "提示"
            dv.prompt = prompt
        dv.showDropDown = False
        ws.add_data_validation(dv)
        dv.add(f"{column_letter}3:{column_letter}{n_rows}")

    add_dv("B", '"每天,每周,每月,每年,隔N天,指定日期"',
           "请选择或输入正确的周期类型", "点击右侧下拉箭头选择，或手动输入", "周期类型")
    add_dv("D", '"周一,周二,周三,周四,周五,周六,周日"',
           "请选择周几", "点击右侧下拉箭头选择周几", "周几(每周)")
    add_dv("E", '"1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,22,23,24,25,26,27,28,29,30,31"',
           "请选择每月的几号", "点击右侧下拉箭头选择几号", "每月几号")
    add_dv("H", '"08:00,09:00,09:30,10:00,10:30,11:00,12:00,12:30,13:00,14:00,14:30,15:00,16:00,17:00,17:30,18:00,19:00,20:00,21:00,22:00"',
           "时间格式应为 时:分，例如 09:30", "点击右侧下拉箭头选择常用时间，或手动输入", "提醒时间")
    add_dv("I", '"强,弱"', "请选择“强”或“弱”", "点击右侧下拉箭头选择：强 / 弱", "提醒强度")

    dv_num = DataValidation(type="whole", operator="between", formula1="1", formula2="3650",
                            allow_blank=True)
    dv_num.errorTitle = "周期数值"
    dv_num.error = "请填写 1~3650 的整数天数"
    dv_num.promptTitle = "周期数值(天)"
    dv_num.prompt = "仅“隔N天”填写，例如 2"
    dv_num.showDropDown = False
    ws.add_data_validation(dv_num)
    dv_num.add(f"C3:C{n_rows}")

    ws.freeze_panes = "A3"
    wb.save(path)


def read_tasks(path):
    """读取 Excel 中的所有提醒任务。返回 {行号: task} 与错误列表。"""
    tasks = {}
    errors = []
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        return tasks, [f"无法打开提醒表格：{e}"]
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        # 定位表头行：新格式(标题+表头在第2行) 或 旧格式(表头在第1行)
        header_row = None
        for hr in (1, 2):
            if len(rows) >= hr and rows[hr - 1] and rows[hr - 1][0] \
                    and "提醒内容" in str(rows[hr - 1][0]):
                header_row = hr
                break
        start_row = (header_row + 1) if header_row else 2
        for idx, row in enumerate(rows, start=1):
            if idx < start_row:
                continue
            if row is None or all(c is None for c in row):
                continue
            content = ""
            if len(row) > 0 and row[0] is not None:
                content = str(row[0]).strip()
            if not content:
                continue

            cycle = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            num = row[2] if len(row) > 2 else None
            weekday_raw = row[3] if len(row) > 3 else None
            dom_raw = row[4] if len(row) > 4 else None
            month_day_raw = row[5] if len(row) > 5 else None
            date_v = row[6] if len(row) > 6 else None
            time_v = row[7] if len(row) > 7 else None
            strength_raw = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""

            if cycle not in ("每天", "每日", "天天", "每周", "每周一次",
                             "每月", "每月一次", "每年", "每年一次",
                             "隔N天", "隔几天", "自定义天数", "间隔",
                             "指定日期", "一次性", "某一天", "单次", "仅一次"):
                errors.append(f"第{idx}行：周期类型“{cycle or '空'}”无法识别，已跳过")
                continue

            task = {
                "content": content,
                "cycle": cycle,
                "num": num,
                "weekday": parse_weekday(weekday_raw),
                "weekday_raw": weekday_raw,
                "day_of_month": parse_day_of_month(dom_raw),
                "dom_raw": dom_raw,
                "month_day": parse_month_day(month_day_raw),
                "month_day_raw": month_day_raw,
                "date": parse_date(date_v),
                "date_raw": date_v,
                "time": parse_time(time_v),
                "time_raw": time_v,
                "strength_raw": strength_raw,
                "strength": strength_raw in ("强", "强提醒", "重要", "紧急", "高"),
            }

            ok = True
            if cycle in ("每周", "每周一次") and task["weekday"] is None:
                errors.append(f"第{idx}行：每周需填写“周几”(D列)，已跳过")
                ok = False
            elif cycle in ("每月", "每月一次") and task["day_of_month"] is None:
                errors.append(f"第{idx}行：每月需填写“几号”(E列)，已跳过")
                ok = False
            elif cycle in ("每年", "每年一次") and task["month_day"] is None:
                errors.append(f"第{idx}行：每年需填写“几月几号”(F列)，已跳过")
                ok = False
            elif cycle in ("隔N天", "隔几天", "自定义天数", "间隔"):
                try:
                    n = int(task["num"])
                    if n < 1:
                        n = 0
                except (TypeError, ValueError):
                    n = 0
                if n < 1:
                    errors.append(f"第{idx}行：隔N天需填写天数(C列)，已跳过")
                    ok = False
                elif task["date"] is None:
                    errors.append(f"第{idx}行：隔N天需填写起始日期(G列)，已跳过")
                    ok = False
            elif cycle in ("指定日期", "一次性", "某一天", "单次", "仅一次") and task["date"] is None:
                errors.append(f"第{idx}行：指定日期需填写具体日期(G列)，已跳过")
                ok = False
            if not ok:
                continue

            task["start_date"] = task["date"] if cycle in ("隔N天", "隔几天", "自定义天数", "间隔") else None
            task["one_date"] = task["date"] if cycle in ("指定日期", "一次性", "某一天", "单次", "仅一次") else None

            tasks[idx] = task
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return tasks, errors


# =====================================================================
# 状态持久化
# =====================================================================
def load_state():
    try:
        with open(STATE_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            return {}
        return data
    except Exception:
        return {}


def save_state(state):
    tmp = STATE_PATH + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
        os.replace(tmp, STATE_PATH)
    except Exception as e:
        print("保存状态失败：", e)


def fingerprint(task):
    raw = "|".join([str(task["content"]), str(task["cycle"]), str(task["num"]),
                    str(task["weekday_raw"]), str(task["dom_raw"]),
                    str(task["month_day_raw"]), str(task["date_raw"]),
                    str(task["time_raw"]), str(task.get("strength_raw"))])
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


# =====================================================================
# 铃铛图标
# =====================================================================
def make_icon_image():
    img = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    # 顶钮（把手）
    d.ellipse((28, 3, 36, 11), fill=(176, 196, 222, 255))
    # 铃身（金色）
    d.polygon([(17, 38), (47, 38), (44, 14), (20, 14)], fill=(240, 180, 41, 255))
    d.ellipse((20, 8, 44, 40), fill=(240, 180, 41, 255))
    # 下沿
    d.ellipse((17, 34, 47, 44), fill=(240, 180, 41, 255))
    d.ellipse((17, 34, 47, 44), outline=(150, 105, 25, 255), width=2)
    # 铃舌（红色）
    d.ellipse((29, 44, 35, 53), fill=(214, 48, 49, 255))
    # 高光
    d.ellipse((25, 13, 33, 20), fill=(255, 214, 122, 255))
    return img


# =====================================================================
# 开机自动启动（写入当前用户注册表 Run 项）
# =====================================================================
RUN_KEY = r"Software\Microsoft\Windows\CurrentVersion\Run"
APP_REG_NAME = "任务提醒"
OLD_REG_NAME = "桌面提醒"


def exe_path_for_startup():
    if getattr(sys, "frozen", False):
        return os.path.abspath(sys.executable)
    pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
    if os.path.exists(pyw):
        return f'"{pyw}" "{os.path.abspath(__file__)}"'
    return os.path.abspath(sys.executable)


def autostart_enabled():
    try:
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, RUN_KEY, 0, winreg.KEY_READ)
        try:
            val, _ = winreg.QueryValueEx(key, APP_REG_NAME)
            return bool(val)
        finally:
            winreg.CloseKey(key)
    except FileNotFoundError:
        return False
    except Exception:
        return False


def set_autostart(enable):
    key = winreg.CreateKey(winreg.HKEY_CURRENT_USER, RUN_KEY)
    try:
        for name in (APP_REG_NAME, OLD_REG_NAME):
            try:
                winreg.DeleteValue(key, name)
            except FileNotFoundError:
                pass
        if enable:
            winreg.SetValueEx(key, APP_REG_NAME, 0, winreg.REG_SZ, exe_path_for_startup())
    finally:
        winreg.CloseKey(key)


# =====================================================================
# 主程序
# =====================================================================
class ReminderApp:
    def __init__(self):
        ensure_default_sounds()
        self.settings = load_settings()

        self.root = tk.Tk()
        self.root.title(APP_NAME)
        self.root.geometry("560x680")
        self.root.resizable(False, False)

        self.tasks = {}
        self.state = load_state()
        self.excel_mtime = None
        self.errors = []

        self.popup = None
        self.current_rid = None
        self.showing = False
        self._shutdown = False
        self._tray = None
        self._flash_job = None
        self._flash_widgets = []
        self._flash_state = 0
        self._flash_enabled = False
        self._flash_colors = ("#ffffff", "#ffffff")
        self.sound_path_vars = {}
        self.autostart_var = tk.BooleanVar(value=autostart_enabled())
        self._player = SoundPlayer()

        first_run = not os.path.exists(STATE_PATH)
        self._build_ui()
        self._load_tasks(force=True)
        self._start_tray()

        self.root.protocol("WM_DELETE_WINDOW", self._hide_to_tray)
        self.root.after(200, self._poll)
        if first_run:
            self.root.after(500, self._ask_autostart_first_run)

    # ---------- 界面 ----------
    def _build_ui(self):
        pad = 12
        top = ttk.Frame(self.root, padding=pad)
        top.pack(fill="both", expand=True)

        title = ttk.Label(top, text=APP_NAME, font=("Microsoft YaHei UI", 18, "bold"))
        title.pack(anchor="w")

        self.status_var = tk.StringVar(value="正在启动……")
        st = ttk.Label(top, textvariable=self.status_var, wraplength=520,
                       justify="left", foreground="#333333")
        st.pack(fill="x", pady=(6, 4))

        ttk.Separator(top, orient="horizontal").pack(fill="x", pady=6)

        ttk.Label(top, text="提醒表格：").pack(anchor="w")
        path_lbl = ttk.Label(top, text=EXCEL_PATH, wraplength=520,
                             foreground="#1a66cc", justify="left")
        path_lbl.pack(fill="x", pady=(0, 6))

        btns = ttk.Frame(top)
        btns.pack(fill="x", pady=4)
        ttk.Button(btns, text="打开提醒表格", command=self._open_excel).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="新建/重置模板", command=self._create_template_click).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="重新读取表格", command=lambda: self._load_tasks(force=True)).pack(side="left")

        ttk.Checkbutton(top, text="开机自动启动（勾选后每次开机自动在后台运行）",
                        variable=self.autostart_var,
                        command=self._toggle_autostart).pack(anchor="w", pady=(8, 0))

        # ---------- 提醒音效设置 ----------
        ttk.Separator(top, orient="horizontal").pack(fill="x", pady=8)
        ttk.Label(top, text="提醒音效：", font=("Microsoft YaHei UI", 11, "bold")).pack(anchor="w")

        self.sound_enable_var = tk.BooleanVar(value=bool(self.settings.get("sound_enabled", True)))
        ttk.Checkbutton(top, text="启用提示音", variable=self.sound_enable_var,
                        command=self._toggle_sound_enabled).pack(anchor="w", pady=(2, 4))

        for label, key in (("强提醒", "sound_strong"), ("弱提醒", "sound_weak"),
                           ("错过提醒", "sound_missed")):
            row = ttk.Frame(top)
            row.pack(fill="x", pady=2)
            ttk.Label(row, text=f"  {label}：", width=8).pack(side="left")
            var = tk.StringVar(value=self._display_sound(key))
            self.sound_path_vars[key] = var
            ttk.Label(row, textvariable=var, foreground="#555555", width=26).pack(side="left")
            ttk.Button(row, text="选择音频…", command=lambda k=key: self._choose_sound(k)).pack(side="left", padx=4)
            ttk.Button(row, text="重置", command=lambda k=key: self._reset_sound(k)).pack(side="left")

        self.err_var = tk.StringVar(value="")
        err = ttk.Label(top, textvariable=self.err_var, wraplength=520,
                        justify="left", foreground="#cc3300")
        err.pack(fill="x", pady=(8, 0))

        bottom = ttk.Frame(top)
        bottom.pack(side="bottom", fill="x")
        ttk.Button(bottom, text="隐藏到托盘", command=self._hide_to_tray).pack(side="left")
        ttk.Button(bottom, text="退出程序", command=self._quit).pack(side="right")

    # ---------- 音效方法 ----------
    def _display_sound(self, key):
        p = self.settings.get(key) or default_sound_path(key)
        return os.path.basename(p) if p else "(使用默认音效)"

    def _toggle_sound_enabled(self):
        self.settings["sound_enabled"] = bool(self.sound_enable_var.get())
        save_settings(self.settings)

    def _choose_sound(self, key):
        path = filedialog.askopenfilename(
            title="选择音效文件",
            filetypes=[("音频文件", "*.wav *.mp3 *.wma *.m4a"), ("所有文件", "*.*")])
        if path:
            self.settings[key] = path
            save_settings(self.settings)
            self.sound_path_vars[key].set(os.path.basename(path))

    def _reset_sound(self, key):
        self.settings[key] = ""
        save_settings(self.settings)
        self.sound_path_vars[key].set("(使用默认音效)")

    def _play_alert(self, style):
        if not self.settings.get("sound_enabled", True):
            return
        key = {"strong": "sound_strong", "weak": "sound_weak", "missed": "sound_missed"}.get(style)
        path = (self.settings.get(key) or "") if key else ""
        if not path:
            path = default_sound_path(style)
        self._player.play_loop(path)   # 循环播放，直到用户点按钮确认

    # ---------- 托盘 ----------
    def _start_tray(self):
        menu = pystray.Menu(
            pystray.MenuItem("显示主窗口", lambda icon, item: self._safe_after(self._show_main)),
            pystray.MenuItem("打开提醒表格", lambda icon, item: self._safe_after(self._open_excel)),
            pystray.MenuItem("重新读取表格", lambda icon, item: self._safe_after(lambda: self._load_tasks(force=True))),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("退出程序", lambda icon, item: self._safe_after(self._quit)),
        )
        icon = pystray.Icon(APP_NAME, make_icon_image(), APP_NAME, menu)
        self._tray = icon
        threading.Thread(target=icon.run, daemon=True).start()

    def _safe_after(self, fn):
        try:
            self.root.after(0, fn)
        except Exception:
            pass

    def _show_main(self):
        self.root.deiconify()
        self.root.lift()
        self.root.attributes("-topmost", True)
        self.root.after(300, lambda: self.root.attributes("-topmost", False))

    def _hide_to_tray(self):
        self.root.withdraw()
        self._notify("已隐藏到托盘", "程序仍在后台运行，会在提醒时间弹窗提示。")

    # ---------- 表格加载 ----------
    def _load_tasks(self, force=False):
        if not os.path.exists(EXCEL_PATH):
            try:
                create_template(EXCEL_PATH)
            except Exception as e:
                self._set_status("创建模板失败：" + str(e))
                return
        try:
            mtime = os.stat(EXCEL_PATH)[stat.ST_MTIME]
        except Exception:
            mtime = None
        if not force and mtime is not None and mtime == self.excel_mtime:
            return

        tasks, errors = read_tasks(EXCEL_PATH)
        self.tasks = tasks
        self.errors = errors

        now = dt.datetime.now()
        new_state = {}
        for rid, task in tasks.items():
            key = str(rid)
            old = self.state.get(key)
            fp = fingerprint(task)
            if old and old.get("fp") == fp and old.get("done") is False:
                entry = dict(old)
                entry["missed"] = False
                nxt = old.get("next")
                if nxt:
                    try:
                        if dt.datetime.fromisoformat(nxt) < now - dt.timedelta(seconds=60):
                            entry["missed"] = True
                    except Exception:
                        pass
                new_state[key] = entry
            else:
                nxt = next_occurrence(task, now)
                new_state[key] = {
                    "fp": fp,
                    "next": nxt.isoformat() if nxt else None,
                    "done": False,
                    "missed": False,
                }
        self.state = new_state
        self.excel_mtime = mtime
        save_state(self.state)

        active = sum(1 for v in self.state.values() if not v.get("done"))
        self._set_status(f"已加载 {active} 条有效提醒（共 {len(tasks)} 行），正在后台监听……")
        if errors:
            self.err_var.set("；".join(errors[:5]))
        else:
            self.err_var.set("")

    def _set_status(self, text):
        self.status_var.set(text)

    def _toggle_autostart(self):
        try:
            set_autostart(self.autostart_var.get())
            if self.autostart_var.get():
                self._set_status("已开启开机自动启动，每次开机都会自动运行")
            else:
                self._set_status("已关闭开机自动启动")
        except Exception as e:
            self.autostart_var.set(autostart_enabled())
            self._set_status("设置开机自启失败：" + str(e))

    def _ask_autostart_first_run(self):
        if messagebox.askyesno(
                "开机自动启动",
                "是否要设置本程序开机自动启动？\n\n"
                "勾选后，每次开机它都会自动在后台运行，并按时弹出提醒。\n"
                "（以后可随时在主窗口取消勾选）"):
            self.autostart_var.set(True)
            self._toggle_autostart()

    def _open_excel(self):
        if not os.path.exists(EXCEL_PATH):
            create_template(EXCEL_PATH)
        try:
            os.startfile(EXCEL_PATH)
        except Exception as e:
            messagebox.showerror("错误", "无法打开表格：" + str(e))

    def _create_template_click(self):
        if os.path.exists(EXCEL_PATH):
            if not messagebox.askyesno("确认",
                                       "将覆盖当前提醒任务表格并重置为示例模板，确定？"):
                return
        create_template(EXCEL_PATH)
        self._load_tasks(force=True)
        self._set_status("已生成模板，正在打开……")
        self._open_excel()

    # ---------- 主循环 ----------
    def _poll(self):
        if self._shutdown:
            return
        try:
            self._load_tasks(force=False)
            self._check_due()
        except Exception as e:
            self._set_status("运行异常：" + str(e))
        self.root.after(5000, self._poll)

    def _check_due(self):
        if self.showing:
            return
        now = dt.datetime.now()
        due = []
        for rid, task in self.tasks.items():
            key = str(rid)
            st = self.state.get(key)
            if not st or st.get("done"):
                continue
            nxt = st.get("next")
            if not nxt:
                continue
            try:
                nxt_dt = dt.datetime.fromisoformat(nxt)
            except Exception:
                continue
            if now >= nxt_dt:
                due.append((rid, task, bool(st.get("missed"))))
        if due:
            self._show_popup(due[0][0], due[0][1], due[0][2])

    # ---------- 弹窗 ----------
    def _flash_stop(self):
        if self._flash_job is not None:
            try:
                if self.popup is not None:
                    self.popup.after_cancel(self._flash_job)
            except Exception:
                pass
            self._flash_job = None

    def _flash_tick(self):
        if not self._flash_enabled or not self.showing or self.popup is None:
            self._flash_job = None
            return
        self._flash_state = 1 - self._flash_state
        on, off = self._flash_colors
        bg = on if self._flash_state else off
        try:
            self.popup.configure(bg=bg)
            for w in self._flash_widgets:
                try:
                    w.configure(bg=bg)
                except Exception:
                    pass
        except Exception:
            pass
        try:
            self._flash_job = self.popup.after(420, self._flash_tick)
        except Exception:
            self._flash_job = None

    def _show_popup(self, rid, task, missed=False):
        self.showing = True
        self.current_rid = rid
        strong = bool(task.get("strength"))

        # 样式：错过=黄色闪烁；强=红色闪烁；弱=蓝色普通
        if missed:
            style = "missed"
        elif strong:
            style = "strong"
        else:
            style = "weak"

        if style == "missed":
            title_suffix = " - 错过的提醒"
            head_text = "⏰ 错过的提醒"
            head_fg = "#b8860b"
            base_bg = "#fffbe6"
            content_font = ("Microsoft YaHei UI", 23, "bold")
            flash_enabled = True
            flash_colors = ("#fff3a8", "#ffffff")
        elif style == "strong":
            title_suffix = " - 重要提醒"
            head_text = "⚠ 重要提醒"
            head_fg = "#cc0000"
            base_bg = "#ffe9e9"
            content_font = ("Microsoft YaHei UI", 24, "bold")
            flash_enabled = True
            flash_colors = ("#ffd6d6", "#ffffff")
        else:
            title_suffix = " - 提醒"
            head_text = "⏰ 到点提醒"
            head_fg = "#1a66cc"
            base_bg = "#eaf3ff"
            content_font = ("Microsoft YaHei UI", 21, "bold")
            flash_enabled = False
            flash_colors = ("#eaf3ff", "#eaf3ff")

        popup = tk.Toplevel(self.root)
        popup.title(APP_NAME + title_suffix)
        popup.attributes("-topmost", True)
        popup.resizable(True, True)
        popup.protocol("WM_DELETE_WINDOW", lambda: self._resolve("ok"))

        self._flash_widgets = []
        frame = tk.Frame(popup, bg=base_bg, padx=26, pady=20)
        frame.pack(fill="both", expand=True)
        self._flash_widgets.append(frame)

        header = tk.Label(frame, text=head_text, bg=base_bg, fg=head_fg,
                          font=("Microsoft YaHei UI", 17, "bold"))
        header.pack(anchor="w")
        self._flash_widgets.append(header)

        info = self._cycle_text(task)
        if missed:
            info = "此前未在电脑前，现补提醒　·　" + info
        info_lbl = tk.Label(frame, text=info, bg=base_bg, fg="#555555",
                            font=("Microsoft YaHei UI", 11), wraplength=620, justify="left")
        info_lbl.pack(anchor="w", pady=(4, 10))
        self._flash_widgets.append(info_lbl)

        # 提醒内容：随内容长度自动加高；过长时自动出现滚动条，保证完整显示
        font_size = content_font[1]
        char_w = max(8, int((620 - 60) / (font_size * 0.62)))
        lines = max(1, math.ceil(len(task["content"]) / char_w))
        text_height = min(max(lines, 2), 12)   # 最多同时显示12行，超出滚动

        content_box = tk.Frame(frame, bg=base_bg)
        content_box.pack(fill="both", expand=True, pady=6)
        content = tk.Text(content_box, bg=base_bg, fg="#111111",
                          font=content_font, wrap="word", height=text_height,
                          width=char_w,
                          relief="flat", highlightthickness=0, padx=4, pady=2)
        content.insert("1.0", task["content"])
        content.configure(state="disabled")
        content.pack(side="left", fill="both", expand=True)
        sb = tk.Scrollbar(content_box, command=content.yview)
        sb.pack(side="right", fill="y")
        content.configure(yscrollcommand=sb.set)
        self._flash_widgets.append(content_box)
        self._flash_widgets.append(content)

        snooze_row = tk.Frame(frame, bg=base_bg)
        snooze_row.pack(fill="x", pady=(30, 6))
        self.snooze_var = tk.IntVar(value=10)
        tk.Label(snooze_row, text="稍后提醒：", bg=base_bg,
                 font=("Microsoft YaHei UI", 13)).pack(side="left")
        spin = tk.Spinbox(snooze_row, from_=1, to=1440, textvariable=self.snooze_var,
                          width=6, font=("Microsoft YaHei UI", 13))
        spin.pack(side="left", padx=(6, 6))
        tk.Label(snooze_row, text=" 分钟", bg=base_bg,
                 font=("Microsoft YaHei UI", 13)).pack(side="left")
        self._flash_widgets.append(snooze_row)

        btns = tk.Frame(frame, bg=base_bg)
        btns.pack(side="bottom", fill="x", pady=(12, 0))
        tk.Button(btns, text="我知道了", command=lambda: self._resolve("ok"),
                  font=("Microsoft YaHei UI", 13, "bold"), width=12).pack(side="left", padx=(0, 10))
        tk.Button(btns, text="稍后提醒", command=lambda: self._resolve("snooze"),
                  font=("Microsoft YaHei UI", 13, "bold"), width=12).pack(side="left", padx=(0, 10))
        tk.Button(btns, text="跳过本次", command=lambda: self._resolve("skip"),
                  font=("Microsoft YaHei UI", 13, "bold"), width=12).pack(side="right")
        self._flash_widgets.append(btns)

        # 按内容自适应窗口大小（居中、限高）
        popup.update_idletasks()
        sw, sh = popup.winfo_screenwidth(), popup.winfo_screenheight()
        req_w = popup.winfo_reqwidth()
        req_h = popup.winfo_reqheight()
        w = min(max(req_w, 520), 700)
        h = min(max(req_h, 330), int(sh * 0.9))
        x = (sw - w) // 2
        y = max(30, (sh - h) // 3)
        popup.geometry(f"{w}x{h}+{x}+{y}")

        self.popup = popup
        self._flash_state = 0
        self._flash_enabled = flash_enabled
        self._flash_colors = flash_colors
        if flash_enabled:
            self._flash_tick()

        # 循环播放对应的提示音（直到用户点按钮确认）
        self._play_alert(style)

        popup.lift()
        popup.focus_force()
        popup.grab_set()

    def _cycle_text(self, task):
        cycle = task["cycle"] or "每天"
        t_str = (task["time"] or DEFAULT_TIME).strftime("%H:%M")
        if cycle in ("隔N天", "隔几天", "自定义天数", "间隔"):
            start = task.get("start_date")
            s = f"周期：每 {task['num']} 天提醒一次（自 {start} 起）{t_str}"
        elif cycle in ("每周", "每周一次"):
            wd = WEEKDAY_CN.get(task.get("weekday"), "?")
            s = f"周期：每周{wd} {t_str} 提醒"
        elif cycle in ("每月", "每月一次"):
            s = f"周期：每月{task['day_of_month']}号 {t_str} 提醒"
        elif cycle in ("每年", "每年一次"):
            mo, da = task["month_day"]
            s = f"周期：每年{mo}月{da}日 {t_str} 提醒"
        elif cycle in ("指定日期", "一次性", "某一天", "单次", "仅一次"):
            s = f"周期：指定日期 {task['one_date']} {t_str} 一次性提醒"
        elif cycle in ("每天", "每日", "天天"):
            s = f"周期：每天 {t_str} 提醒"
        else:
            s = f"周期：每天 {t_str} 提醒"
        s += "　|　" + ("强提醒" if task.get("strength") else "弱提醒")
        return s

    def _resolve(self, action):
        self._flash_stop()
        self._player.stop()   # 用户已处理，停止循环提示音
        if self.popup is not None:
            try:
                self.popup.destroy()
            except Exception:
                pass
        self.popup = None
        rid = self.current_rid
        self.current_rid = None
        self.showing = False

        key = str(rid)
        st = self.state.get(key)
        task = self.tasks.get(rid)
        if st is None or task is None:
            self.root.after(300, self._check_due)
            return

        now = dt.datetime.now()
        one_time = task["cycle"] in ("指定日期", "一次性", "某一天", "单次", "仅一次")

        if action == "ok":
            if one_time:
                st["done"] = True
                st["next"] = None
            else:
                nxt = next_occurrence(task, now, strict=True)
                st["next"] = nxt.isoformat() if nxt else None
        elif action == "skip":
            nxt = next_occurrence(task, now, strict=True)
            st["next"] = nxt.isoformat() if nxt else None
        elif action == "snooze":
            try:
                minutes = int(self.snooze_var.get())
            except Exception:
                minutes = 10
            if minutes < 1:
                minutes = 1
            if minutes > 1440:
                minutes = 1440
            st["next"] = (now + dt.timedelta(minutes=minutes)).isoformat()

        st["missed"] = False
        save_state(self.state)
        self.root.after(300, self._check_due)

    # ---------- 通知 / 退出 ----------
    def _notify(self, title, msg):
        try:
            import ctypes
            ctypes.windll.user32.MessageBoxW(0, msg, title, 0x40)
        except Exception:
            pass

    def _quit(self):
        self._shutdown = True
        self._player.stop()
        try:
            if self._tray is not None:
                self._tray.stop()
        except Exception:
            pass
        try:
            self.root.destroy()
        except Exception:
            pass


def main():
    block_network()   # 代码层面禁止联网

    import ctypes
    mutex = ctypes.windll.kernel32.CreateMutexW(None, False, "Global\\DoubaoTaskReminderMutex")
    if ctypes.windll.kernel32.GetLastError() == 183:
        try:
            ctypes.windll.user32.MessageBoxW(0, "任务提醒已在运行，请查看系统托盘。", APP_NAME, 0x40)
        except Exception:
            pass
        return

    app = ReminderApp()
    app.root.mainloop()


if __name__ == "__main__":
    main()
